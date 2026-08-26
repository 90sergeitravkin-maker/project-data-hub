# src/back/app_data_validator/validator.py
"""
DataValidator — проверка CSV/Parquet/XLSX/XLS с использованием DuckDB для массовой обработки.
Автоматически ищет последнюю доступную версию справочника по дате в имени папки.
Поддерживает валидацию как отдельных файлов, так и целых директорий.
"""
import csv
import re
import time
import duckdb
from datetime import datetime, date
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple
from src.core.logger import logger
from src.back.app_data_validator.config import BASE_DATA_DIR

try:
    import pyarrow.parquet as pq
except ImportError:
    pq = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    import xlrd
except ImportError:
    xlrd = None


class DataValidator:
    """
    Валидатор файлов данных с массовой обработкой через DuckDB.
    """
    SUPPORTED_EXTENSIONS = {'csv', 'parquet', 'xlsx', 'xls'}
    SUPPORTED_EXTENSIONS_DISPLAY = ['CSV', 'Parquet', 'XLSX', 'XLS']

    def __init__(self, config: dict, max_error_examples: int = 1000, max_duplicate_examples: int = 10):
        self.config = config
        self.max_error_examples = max_error_examples
        self.max_duplicate_examples = max_duplicate_examples
        self._ref_cache: Dict[Tuple[str, str], Set[str]] = {}

    # =========================================================================
    # 1. ПОИСК СПРАВОЧНИКА С МАКСИМАЛЬНОЙ ДАТОЙ
    # =========================================================================
    def _find_reference_file_with_max_date(self, reference_name: str) -> Optional[str]:
        """
        Находит файл справочника в папке с максимальной датой.
        """
        logger.debug(f"[Validator] Поиск справочника: {reference_name}")
        ref_dir = BASE_DATA_DIR / reference_name
        if not ref_dir.exists() or not ref_dir.is_dir():
            logger.warning(f"[Validator] Папка справочника не найдена: {ref_dir}")
            return None

        date_pattern = re.compile(r'^\d{4}-\d{2}-\d{2}$')
        date_dirs = []
        for item in ref_dir.iterdir():
            if item.is_dir() and date_pattern.match(item.name):
                try:
                    datetime.strptime(item.name, "%Y-%m-%d")
                    date_dirs.append(item)
                except ValueError:
                    pass

        if not date_dirs:
            logger.warning(f"[Validator] Папок с датами не найдено в {ref_dir}")
            return self._find_file_in_directory(ref_dir)

        date_dirs.sort(key=lambda x: x.name, reverse=True)
        latest_date_dir = date_dirs[0]
        logger.info(f"[Validator] Найдена папка с максимальной датой: {latest_date_dir.name}")

        file_path = self._find_file_in_directory(latest_date_dir)
        if file_path:
            logger.info(f"[Validator] Найден файл справочника: {file_path}")
            return file_path

        for date_dir in date_dirs[1:]:
            file_path = self._find_file_in_directory(date_dir)
            if file_path:
                logger.info(f"[Validator] Найден файл в папке {date_dir.name}: {file_path}")
                return file_path

        logger.warning(f"[Validator] Файлы не найдены в папках с датами: {ref_dir}")
        return None

    def _find_file_in_directory(self, directory: Path) -> Optional[str]:
        """Ищет файл данных в директории. Приоритет: parquet > csv > xlsx > xls"""
        for ext in ['.parquet', '.csv', '.xlsx', '.xls']:
            files = list(directory.glob(f"*{ext}"))
            if files:
                files.sort(key=lambda x: x.name)
                return str(files[0])
        return None

    # =========================================================================
    # 2. ЗАГРУЗКА ЗНАЧЕНИЙ ИЗ СПРАВОЧНИКА
    # =========================================================================
    def _load_reference_set(self, reference_name: str, column_name: str) -> Set[str]:
        key = (reference_name, column_name)
        if key in self._ref_cache:
            logger.debug(f"[Validator] Справочник из кэша: {reference_name}.{column_name}")
            return self._ref_cache[key]

        logger.debug(f"[Validator] Загрузка справочника: {reference_name}.{column_name}")
        file_path = self._find_reference_file_with_max_date(reference_name)
        if not file_path:
            logger.error(f"[Validator] Справочник не найден: {reference_name}")
            self._ref_cache[key] = set()
            return set()

        ext = Path(file_path).suffix.lower().lstrip('.')
        file_size = Path(file_path).stat().st_size
        if file_size == 0:
            logger.error(f"[Validator] Справочник пуст: {file_path}")
            self._ref_cache[key] = set()
            return set()

        logger.info(f"[Validator] Загрузка справочника: {file_path}, размер: {file_size} байт")
        start_time = time.time()
        values: Set[str] = set()

        try:
            safe_path = file_path.replace("'", "''")
            with duckdb.connect() as conn:
                if ext == 'parquet':
                    table_expr = f"read_parquet('{safe_path}')"
                elif ext == 'csv':
                    table_expr = f"read_csv_auto('{safe_path}')"
                elif ext in ('xlsx', 'xls'):
                    values = self._read_excel_fallback(file_path, column_name)
                    elapsed_ms = round((time.time() - start_time) * 1000, 2)
                    logger.info(f"[Validator] Справочник загружен: {len(values)} записей, elapsed_ms={elapsed_ms}")
                    self._ref_cache[key] = values
                    return values
                else:
                    logger.warning(f"[Validator] Неподдерживаемый формат: {ext}")
                    self._ref_cache[key] = set()
                    return set()

                try:
                    describe_query = f"DESCRIBE SELECT * FROM {table_expr}"
                    columns = conn.execute(describe_query).fetchall()
                    column_names = [c[0] for c in columns]
                    if column_name not in column_names:
                        logger.error(
                            f"[Validator] Колонка '{column_name}' не найдена. Доступные: {column_names[:10]}...")
                        self._ref_cache[key] = set()
                        return set()
                except Exception as e:
                    logger.warning(f"[Validator] Не удалось проверить схему: {e}")

                query = f"""
                    SELECT DISTINCT CAST("{column_name}" AS VARCHAR) 
                    FROM {table_expr} 
                    WHERE "{column_name}" IS NOT NULL 
                    AND TRIM(CAST("{column_name}" AS VARCHAR)) != ''
                """
                result = conn.execute(query).fetchall()
                for row in result:
                    if row[0] is not None and str(row[0]).strip():
                        values.add(str(row[0]).strip())
        except Exception as e:
            logger.error(f"[Validator] Ошибка загрузки справочника {file_path}: {e}", exc_info=True)

        elapsed_ms = round((time.time() - start_time) * 1000, 2)
        if len(values) == 0:
            logger.warning(f"[Validator] В справочнике не найдено значений: {reference_name}.{column_name}")
        else:
            logger.info(
                f"[Validator] Справочник загружен: {len(values)} записей, примеры: {list(values)[:5]}, elapsed_ms={elapsed_ms}")

        self._ref_cache[key] = values
        return values

    def _read_excel_fallback(self, file_path: str, column_name: str) -> Set[str]:
        values: Set[str] = set()
        ext = Path(file_path).suffix.lower()
        try:
            if ext == '.xlsx' and load_workbook:
                wb = load_workbook(file_path, read_only=True)
                ws = wb.active
                header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
                try:
                    col_idx = header.index(column_name)
                except ValueError:
                    wb.close()
                    return values
                for row in ws.iter_rows(min_row=2, values_only=True):
                    v = row[col_idx] if col_idx < len(row) else None
                    if v is not None and str(v).strip():
                        values.add(str(v).strip())
                wb.close()
            elif ext == '.xls' and xlrd:
                book = xlrd.open_workbook(file_path, on_demand=True)
                sheet = book.sheet_by_index(0)
                header = sheet.row_values(0)
                try:
                    col_idx = header.index(column_name)
                except ValueError:
                    return values
                for row_idx in range(1, sheet.nrows):
                    v = sheet.cell_value(row_idx, col_idx)
                    if v is not None and str(v).strip():
                        values.add(str(v).strip())
        except Exception as e:
            logger.error(f"[Validator] Ошибка чтения Excel {file_path}: {e}")
        return values

    def _preload_reference_sets(self, column_rules: Dict) -> Dict[Tuple[str, str], Set[str]]:
        ref_sets = {}
        for col, rules in column_rules.items():
            for check in rules.get('checks', []):
                if 'reference_file' in check and 'reference_column' in check:
                    key = (check['reference_file'], check['reference_column'])
                    if key not in ref_sets:
                        ref_sets[key] = self._load_reference_set(check['reference_file'], check['reference_column'])
        return ref_sets

    # =========================================================================
    # 3. ОСНОВНАЯ ВАЛИДАЦИЯ (Файл или Директория)
    # =========================================================================
    def _extract_source_from_path(self, file_path: str) -> Optional[str]:
        """
        Извлекает имя источника из начала пути.
        Ищет совпадение с ключами self.config.
        """
        normalized_path = file_path.replace('\\', '/').strip('/')
        parts = normalized_path.split('/')
        if not parts:
            return None

        first_part = parts[0]
        if first_part in self.config:
            return first_part

        first_part_lower = first_part.lower()
        for source in self.config.keys():
            if source.lower() == first_part_lower:
                return source

        for source in self.config.keys():
            if source in normalized_path or normalized_path in source:
                return source

        return None

    def validate_file(self, file_path: str, source_name: str) -> Dict[str, Any]:
        """
        Валидация файла или директории.
        """
        start_time = time.time()
        logger.info(f"[Validator] Начало валидации: source={source_name}, path={file_path}")

        if source_name not in self.config:
            return {"error": f"Источник {source_name} не найден в конфигурации"}

        column_rules = self.config[source_name]
        resolved_path = self._resolve_file_path(file_path)
        path_obj = Path(resolved_path)

        if not path_obj.exists():
            return {
                "error": f"Путь не существует: {file_path}",
                "full_path": str(path_obj),
                "hint": f"Проверьте, что папка или файл существуют. "
                        f"Используйте абсолютный путь или убедитесь, что путь указан "
                        f"относительно BASE_DATA_DIR ({BASE_DATA_DIR})"
            }

        # ============================================================
        # ОБРАБОТКА: если путь указывает на директорию (ИСПРАВЛЕНО: rglob)
        # ============================================================
        if path_obj.is_dir():
            supported_exts = {f'.{ext}' for ext in self.SUPPORTED_EXTENSIONS}

            # ИСПРАВЛЕНИЕ: используем rglob для рекурсивного поиска во всех подпапках
            files_found = [
                f for f in path_obj.rglob('*')
                if f.is_file() and f.suffix.lower() in supported_exts
            ]

            if not files_found:
                # Также используем rglob, чтобы честно проверить всё дерево папок
                all_files = [f for f in path_obj.rglob('*') if f.is_file()]
                if all_files:
                    extensions = sorted({f.suffix.lower() for f in all_files})
                    return {
                        "error": f"В директории {file_path} и её подпапках нет поддерживаемых файлов",
                        "full_path": str(path_obj),
                        "supported_extensions": list(self.SUPPORTED_EXTENSIONS),
                        "found_extensions": extensions,
                        "hint": f"Найдены файлы с расширениями: {extensions}. Поддерживаются: {list(self.SUPPORTED_EXTENSIONS)}"
                    }
                else:
                    return {
                        "error": f"Директория {file_path} и все её подпапки пусты",
                        "full_path": str(path_obj),
                        "hint": "В папке нет файлов. Проверьте, что данные загружены."
                    }

            return self._validate_directory(path_obj, source_name, column_rules, start_time)

        # ============================================================
        # ОБРАБОТКА: если путь указывает на файл
        # ============================================================
        ext = path_obj.suffix.lower().lstrip('.')
        if ext not in self.SUPPORTED_EXTENSIONS:
            return {
                "error": f"Неподдерживаемый формат файла: {ext or 'без расширения'}",
                "full_path": str(path_obj),
                "supported_extensions": list(self.SUPPORTED_EXTENSIONS),
                "hint": f"Файл {path_obj.name} имеет расширение '{ext}', но поддерживаются только: {', '.join(self.SUPPORTED_EXTENSIONS_DISPLAY)}"
            }

        actual_columns = self._get_columns(resolved_path, ext)
        if actual_columns is None:
            return {
                "error": f"Не удалось прочитать заголовки файла: {path_obj.name}",
                "full_path": str(path_obj),
                "hint": "Возможно, файл повреждён или имеет нестандартный формат."
            }

        required_cols = [c for c, r in column_rules.items() if r.get('required', False)]
        missing = [c for c in required_cols if c not in actual_columns]
        if missing:
            suggestions = self._find_similar_columns(missing, actual_columns)
            return {
                "error": f"Отсутствуют обязательные колонки: {missing}",
                "full_path": str(path_obj),
                "available_columns": actual_columns,
                "suggestions": suggestions,
                "hint": f"В файле {path_obj.name} отсутствуют колонки: {missing}. "
                        f"Возможно, вы используете не тот source_name."
            }

        if ext in ('parquet', 'csv'):
            result = self._validate_with_duckdb(resolved_path, ext, column_rules, start_time)
        else:
            result = self._validate_excel(resolved_path, ext, column_rules, start_time)

        result["source"] = source_name
        return result

    def _find_similar_columns(self, missing: List[str], available: List[str]) -> Dict[str, List[str]]:
        suggestions = {}
        available_lower = {col.lower(): col for col in available}
        for col in missing:
            col_lower = col.lower()
            similar = []
            if col_lower in available_lower:
                similar.append(f"ТОЧНОЕ СОВПАДЕНИЕ (регистр): {available_lower[col_lower]}")
            for avail in available:
                avail_lower = avail.lower()
                if col_lower in avail_lower or avail_lower in col_lower:
                    if avail not in similar:
                        similar.append(avail)
                col_normalized = col_lower.replace('_', ' ').replace('-', ' ')
                avail_normalized = avail_lower.replace('_', ' ').replace('-', ' ')
                if col_normalized == avail_normalized and avail not in similar:
                    similar.append(avail)
            suggestions[col] = similar[:5] if similar else []
        return suggestions

    # =========================================================================
    # 4. ВСПОМОГАТЕЛЬНЫЕ МЕТОДЫ
    # =========================================================================
    @staticmethod
    def _resolve_file_path(file_path: str) -> str:
        """Преобразует путь в абсолютный с проверкой структуры."""
        path = Path(file_path)
        if path.is_absolute():
            if path.exists():
                return str(path)

        absolute_path = BASE_DATA_DIR / file_path
        parts = Path(file_path).parts
        if len(parts) >= 3:
            entity_path = BASE_DATA_DIR / parts[0]
            if entity_path.exists() and entity_path.is_dir():
                date_path = entity_path / parts[1]
                if date_path.exists() and date_path.is_dir():
                    period_path = date_path / parts[2]
                    if period_path.exists() and period_path.is_dir():
                        return str(period_path)
                    for sub in date_path.iterdir():
                        if sub.is_dir() and sub.name.startswith(parts[2][:4]):
                            return str(sub)

        if absolute_path.exists():
            return str(absolute_path)
        return str(absolute_path)

    def _validate_directory(self, dir_path: Path, source_name: str, column_rules: Dict, start_time: float) -> Dict[
        str, Any]:
        logger.info(f"[Validator] Валидация директории (рекурсивно): {dir_path}")
        supported_exts = {'.csv', '.parquet', '.xlsx', '.xls'}

        # ИСПРАВЛЕНИЕ: rglob вместо iterdir для обхода всех вложенных папок
        files_to_process = [
            f for f in dir_path.rglob('*')
            if f.is_file() and f.suffix.lower() in supported_exts
        ]

        if not files_to_process:
            return {"error": f"В директории {dir_path} не найдено файлов с расширениями {supported_exts}"}

        total_rows = 0
        all_errors = []
        error_summary = {}
        files_processed = []
        first_file_columns = None
        required_cols = [c for c, r in column_rules.items() if r.get('required', False)]

        for file_path in files_to_process:
            ext = file_path.suffix.lower().lstrip('.')
            logger.info(f"[Validator] Обработка файла: {file_path.name}")

            actual_columns = self._get_columns(str(file_path), ext)
            if first_file_columns is None and actual_columns:
                first_file_columns = actual_columns

            missing = [c for c in required_cols if c not in actual_columns]
            if missing:
                all_errors.append({
                    "category": "schema_mismatch",
                    "column_name": ", ".join(missing),
                    "value": file_path.name,
                    "error_text": f"В файлах отсутствуют обязательные колонки: {missing}",
                    "count": len(files_to_process),
                    "row_num": None,
                    "file_name": file_path.name
                })
                error_summary["schema_mismatch"] = len(files_to_process)
                logger.error(
                    f"[Validator] СХЕМА НЕ СОВПАДАЕТ! "
                    f"Обязательные колонки: {required_cols}, "
                    f"Доступные: {actual_columns[:20]}..."
                )
                return {
                    "file": str(dir_path),
                    "source": source_name,
                    "total_rows": 0,
                    "status": "FAIL",
                    "error_summary": error_summary,
                    "errors": all_errors,
                    "files_processed": [],
                    "diagnostic": {
                        "required_columns": required_cols,
                        "actual_columns_sample": actual_columns[:20],
                        "missing_columns": missing,
                        "hint": f"Схема файлов не соответствует ожидаемой. "
                                f"Проверьте, что вы используете правильный source_name. "
                                f"Доступные источники: {list(self.config.keys())}"
                    }
                }

            if actual_columns is None:
                all_errors.append({
                    "category": "file_processing_error",
                    "column_name": None,
                    "value": file_path.name,
                    "error_text": "Не удалось прочитать заголовки",
                    "count": 1,
                    "row_num": None,
                    "file_name": file_path.name
                })
                error_summary["file_processing_error"] = error_summary.get("file_processing_error", 0) + 1
                continue

            missing = [c for c in required_cols if c not in actual_columns]
            if missing:
                all_errors.append({
                    "category": "missing_columns",
                    "column_name": ", ".join(missing),
                    "value": file_path.name,
                    "error_text": f"Отсутствуют обязательные колонки: {missing}",
                    "count": 1,
                    "row_num": None,
                    "file_name": file_path.name
                })
                error_summary["missing_columns"] = error_summary.get("missing_columns", 0) + 1
                continue

            if ext in ('parquet', 'csv'):
                result = self._validate_with_duckdb(str(file_path), ext, column_rules, start_time)
            else:
                result = self._validate_excel(str(file_path), ext, column_rules, start_time)

            if "error" in result:
                all_errors.append({
                    "category": "file_processing_error",
                    "column_name": None,
                    "value": file_path.name,
                    "error_text": result["error"],
                    "count": 1,
                    "row_num": None,
                    "file_name": file_path.name
                })
                error_summary["file_processing_error"] = error_summary.get("file_processing_error", 0) + 1
                continue

            total_rows += result.get("total_rows", 0)
            files_processed.append(str(file_path.relative_to(dir_path)))  # Относительный путь для удобства

            for err in result.get("errors", []):
                err["file_name"] = file_path.name
                all_errors.append(err)
                cat = err["category"]
                error_summary[cat] = error_summary.get(cat, 0) + err.get("count", 1)

        all_errors.sort(key=lambda x: x.get("count", 1), reverse=True)
        elapsed_ms = round((time.time() - start_time) * 1000, 2)
        status = "OK" if not all_errors else "FAIL"

        logger.info(
            f"[Validator] Завершена валидация директории: status={status}, files={len(files_processed)}, total_rows={total_rows}, errors={len(all_errors)}, elapsed_ms={elapsed_ms}")

        result = {
            "file": str(dir_path),
            "source": source_name,
            "total_rows": total_rows,
            "status": status,
            "error_summary": error_summary,
            "errors": all_errors[:self.max_error_examples],
            "files_processed": files_processed,
        }

        if status == "FAIL" and first_file_columns:
            result["diagnostic"] = {
                "required_columns": required_cols,
                "actual_columns_sample": first_file_columns[:20],
                "total_columns_found": len(first_file_columns),
                "columns_match": [c for c in required_cols if c in first_file_columns],
                "columns_missing": [c for c in required_cols if c not in first_file_columns],
                "hint": "Проверьте, что вы используете правильный source_name. "
                        f"Доступные источники: {list(self.config.keys())}"
            }

        return result

    # =========================================================================
    # 5. МАССОВАЯ ВАЛИДАЦИЯ ЧЕРЕЗ DUCKDB
    # =========================================================================
    def _validate_with_duckdb(self, file_path: str, ext: str, column_rules: Dict, start_time: float) -> Dict[str, Any]:
        safe_path = file_path.replace("'", "''")
        table_expr = f"read_parquet('{safe_path}')" if ext == 'parquet' else f"read_csv_auto('{safe_path}')"

        with duckdb.connect() as conn:
            total_rows = conn.execute(f"SELECT COUNT(*) FROM {table_expr}").fetchone()[0]
            if total_rows == 0:
                return {"file": file_path, "total_rows": 0, "status": "OK", "error_summary": {}, "errors": []}

            ref_sets = self._preload_reference_sets(column_rules)
            ref_tables = {}
            for (ref_name, ref_col), ref_set in ref_sets.items():
                if ref_set:
                    table_name = f"ref_{abs(hash(ref_name))}_{abs(hash(ref_col))}".replace('-', '_')
                    conn.execute(f"CREATE TEMPORARY TABLE {table_name} (value VARCHAR)")
                    conn.executemany(f"INSERT INTO {table_name} VALUES (?)", [(v,) for v in ref_set])
                    ref_tables[(ref_name, ref_col)] = table_name

            error_queries = []
            query_errors = []

            for col, rules in column_rules.items():
                col_safe = col.replace('"', '""')
                col_quoted = f'"{col_safe}"'

                if rules.get('required', False):
                    error_queries.append(f"""
                        SELECT 'required_null' as category, '{col}' as column_name, NULL as value, 
                               'Поле не может быть NULL' as error_text, COUNT(*) as count, MIN(row_num) as first_row
                        FROM (SELECT row_number() OVER () as row_num, * FROM {table_expr}) t
                        WHERE {col_quoted} IS NULL OR TRIM(CAST({col_quoted} AS VARCHAR)) = ''
                        HAVING COUNT(*) > 0
                    """)

                for check in rules.get('checks', []):
                    if 'reference_file' in check and 'reference_column' in check:
                        ref_key = (check['reference_file'], check['reference_column'])
                        ref_table = ref_tables.get(ref_key)
                        if ref_table:
                            error_queries.append(f"""
                                SELECT 'reference_failed' as category, '{col}' as column_name, CAST({col_quoted} AS VARCHAR) as value, 
                                       'Значение не найдено в справочнике' as error_text, COUNT(*) as count, MIN(row_num) as first_row
                                FROM (SELECT row_number() OVER () as row_num, * FROM {table_expr}) t
                                WHERE {col_quoted} IS NOT NULL AND TRIM(CAST({col_quoted} AS VARCHAR)) != ''
                                  AND CAST({col_quoted} AS VARCHAR) NOT IN (SELECT value FROM {ref_table})
                                GROUP BY CAST({col_quoted} AS VARCHAR) HAVING COUNT(*) > 0
                            """)
                        else:
                            error_queries.append(f"""
                                SELECT 'reference_not_loaded' as category, '{col}' as column_name, CAST({col_quoted} AS VARCHAR) as value, 
                                       'Справочник не загружен' as error_text, COUNT(*) as count, MIN(row_num) as first_row
                                FROM (SELECT row_number() OVER () as row_num, * FROM {table_expr}) t
                                WHERE {col_quoted} IS NOT NULL AND TRIM(CAST({col_quoted} AS VARCHAR)) != ''
                                GROUP BY CAST({col_quoted} AS VARCHAR) HAVING COUNT(*) > 0
                            """)

                if rules.get('unique', False):
                    error_queries.append(f"""
                        WITH ranked AS (
                            SELECT {col_quoted} as value, row_number() OVER () as row_num, COUNT(*) OVER (PARTITION BY {col_quoted}) as cnt
                            FROM {table_expr} WHERE {col_quoted} IS NOT NULL AND TRIM(CAST({col_quoted} AS VARCHAR)) != ''
                        )
                        SELECT 'duplicate' as category, '{col}' as column_name, CAST(value AS VARCHAR) as value, 
                               'Дубликат' as error_text, cnt as count, MIN(row_num) as first_row
                        FROM ranked WHERE cnt > 1 GROUP BY value, cnt LIMIT {self.max_duplicate_examples}
                    """)

            all_errors = []
            for query in error_queries:
                try:
                    for row in conn.execute(query).fetchall():
                        all_errors.append({
                            "category": row[0],
                            "column_name": row[1],
                            "value": row[2],
                            "error_text": row[3],
                            "count": row[4],
                            "row_num": row[5]
                        })
                except Exception as e:
                    error_msg = str(e)
                    logger.warning(f"[Validator] Ошибка выполнения запроса: {error_msg}")

                    column_name = "unknown"
                    match = re.search(r'column "([^"]+)"', error_msg, re.IGNORECASE)
                    if match:
                        column_name = match.group(1)
                    else:
                        match = re.search(r'Column "([^"]+)"', error_msg)
                        if match:
                            column_name = match.group(1)

                    short_error = error_msg
                    if '\n' in short_error:
                        short_error = short_error.split('\n')[0]
                    if 'LINE 4:' in short_error:
                        short_error = short_error.split('LINE 4:')[0].strip()
                    if len(short_error) > 200:
                        short_error = short_error[:200] + "..."

                    query_errors.append({
                        "category": "query_execution_error",
                        "column_name": column_name,
                        "value": None,
                        "error_text": short_error,
                        "count": 1,
                        "row_num": None
                    })

            if query_errors:
                all_errors.extend(query_errors)

            error_summary = {}
            for err in all_errors:
                error_summary[err["category"]] = error_summary.get(err["category"], 0) + err["count"]

            all_errors.sort(key=lambda x: x["count"], reverse=True)
            has_errors = bool(all_errors)
            status = "FAIL" if has_errors else "OK"

            return {
                "file": file_path,
                "total_rows": total_rows,
                "status": status,
                "error_summary": error_summary,
                "errors": all_errors[:self.max_error_examples],
            }

    # =========================================================================
    # 6. ВАЛИДАЦИЯ EXCEL (построчная)
    # =========================================================================
    def _validate_excel(self, file_path: str, ext: str, column_rules: Dict, start_time: float) -> Dict[str, Any]:
        errors, error_summary, total_rows = [], {}, 0
        try:
            ref_sets = self._preload_reference_sets(column_rules)
            if ext == 'xlsx' and load_workbook:
                wb = load_workbook(file_path, read_only=True)
                ws = wb.active
                header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
                col_index = {h: i for i, h in enumerate(header) if h is not None}
                for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
                    total_rows += 1
                    self._process_row_excel(
                        {col: row[col_index[col]] if col in col_index else None for col in column_rules}, row_num,
                        column_rules, ref_sets, errors, error_summary)
                wb.close()
            elif ext == 'xls' and xlrd:
                book = xlrd.open_workbook(file_path, on_demand=True)
                sheet = book.sheet_by_index(0)
                header = sheet.row_values(0)
                col_index = {h: i for i, h in enumerate(header) if h}
                for row_num in range(1, sheet.nrows):
                    total_rows += 1
                    self._process_row_excel(
                        {col: sheet.cell_value(row_num, col_index[col]) if col in col_index else None for col in
                         column_rules},
                        row_num, column_rules, ref_sets, errors, error_summary)
            else:
                return {"error": f"Библиотека для чтения {ext} не установлена"}
        except Exception as e:
            logger.error(f"[Validator] Ошибка валидации Excel {file_path}: {e}", exc_info=True)
            return {"error": f"Ошибка валидации Excel: {e}"}

        return {
            "file": file_path,
            "total_rows": total_rows,
            "status": "OK" if not errors else "FAIL",
            "error_summary": error_summary,
            "errors": errors[:self.max_error_examples],
        }

    def _process_row_excel(self, row: Dict, row_num: int, column_rules: Dict, ref_sets: Dict, errors: List,
                           error_summary: Dict) -> None:
        for col, rules in column_rules.items():
            value = row.get(col)
            value_str = str(value).strip() if value is not None else ""

            if rules.get('required', False) and (value is None or value_str == ''):
                errors.append({"category": "required_null", "column_name": col, "value": None,
                               "error_text": "Поле не может быть NULL", "count": 1, "row_num": row_num})
                error_summary["required_null"] = error_summary.get("required_null", 0) + 1

            for check in rules.get('checks', []):
                if 'reference_file' not in check or 'reference_column' not in check:
                    continue
                ref_key = (check['reference_file'], check['reference_column'])
                ref_set = ref_sets.get(ref_key)

                if value is not None and value_str:
                    if ref_set is None:
                        errors.append({"category": "reference_not_loaded", "column_name": col, "value": value_str,
                                       "error_text": "Справочник не загружен", "count": 1, "row_num": row_num})
                        error_summary["reference_not_loaded"] = error_summary.get("reference_not_loaded", 0) + 1
                    elif value_str not in ref_set:
                        errors.append({"category": "reference_failed", "column_name": col, "value": value_str,
                                       "error_text": f"Значение не найдено в справочнике", "count": 1,
                                       "row_num": row_num})
                        error_summary["reference_failed"] = error_summary.get("reference_failed", 0) + 1
                elif check.get('is_null') is False:
                    errors.append({"category": "required_null", "column_name": col, "value": None,
                                   "error_text": "Поле не может быть NULL", "count": 1, "row_num": row_num})
                    error_summary["required_null"] = error_summary.get("required_null", 0) + 1

    # =========================================================================
    # 7. ВСПОМОГАТЕЛЬНЫЕ МЕТОДЫ
    # =========================================================================
    def _get_columns(self, file_path: str, ext: str) -> Optional[List[str]]:
        try:
            if ext == 'csv':
                with open(file_path, 'r', encoding='utf-8') as f:
                    return [h.strip() for h in next(csv.reader(f), [])]
            if ext == 'parquet' and pq:
                return list(pq.read_schema(file_path).names)
            if ext == 'xlsx' and load_workbook:
                wb = load_workbook(file_path, read_only=True)
                ws = wb.active
                columns = [str(cell).strip() if cell else '' for cell in
                           next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
                wb.close()
                return columns
            if ext == 'xls' and xlrd:
                book = xlrd.open_workbook(file_path, on_demand=True)
                return [str(cell).strip() for cell in book.sheet_by_index(0).row_values(0)]
        except Exception as e:
            logger.error(f"[Validator] Ошибка чтения заголовков {file_path}: {e}", exc_info=True)
            return None
